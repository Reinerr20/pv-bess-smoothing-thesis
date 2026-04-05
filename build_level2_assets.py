from __future__ import annotations

from pathlib import Path
import re

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.lines import Line2D
from matplotlib.transforms import blended_transform_factory

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# =========================
# 0) PATHS
# =========================
FEAS_SUMMARY_CSV = Path("outputs_level2_feasibility") / "level2_feasibility_results.csv"
FEAS_PROFILE_DIR = Path("outputs_level2_feasibility") / "feasibility_profiles"

OUT_DIR = Path("outputs_level2_assets")
FIG_DIR = OUT_DIR / "figures"
TAB_DIR = OUT_DIR / "tables_word"

FIG_DIR.mkdir(parents=True, exist_ok=True)
TAB_DIR.mkdir(parents=True, exist_ok=True)

# fixed representative mechanism case
MECH_LABEL = "cloudy"
MECH_DATE = "2024-04-03"
MECH_WINDOW = "60min"
MECH_SCENARIO = "small"

PROFILE_RE = re.compile(
    r"^(SMA|EMA)_(clear|cloudy|medium)_(\d{4}-\d{2}-\d{2})_(\d+min)_(small|large)\.csv$",
    re.IGNORECASE,
)

LABEL_ORDER = {"clear": 0, "medium": 1, "cloudy": 2}
METHOD_ORDER = {"SMA": 0, "EMA": 1}
SCENARIO_ORDER = {"small": 0, "large": 1}
WINDOW_ORDER = {"10min": 10, "20min": 20, "30min": 30, "60min": 60}

COLORS_DAY = {
    "clear": "tab:blue",
    "medium": "tab:orange",
    "cloudy": "tab:green",
}
MARKERS_METHOD = {
    "SMA": "o",
    "EMA": "^",
}
LINESTYLE_SCEN = {
    "small": "-",
    "large": "--",
}


# =========================
# 1) LOADERS
# =========================
def load_feasibility_summary(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Feasibility summary not found: {path}")

    df = pd.read_csv(path)

    required = [
        "label", "date", "method", "window_name", "scenario",
        "ramp_p95_kw_target", "ramp_p95_kw_actual",
        "throughput_req_kwh", "throughput_act_kwh",
        "energy_mismatch_kwh", "power_limit_frac", "soc_limit_frac"
    ]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns in feasibility summary: {missing}")

    df["label_order"] = df["label"].map(LABEL_ORDER)
    df["method_order"] = df["method"].map(METHOD_ORDER)
    df["scenario_order"] = df["scenario"].map(SCENARIO_ORDER)
    df["window_order"] = df["window_name"].map(WINDOW_ORDER)

    df = df.sort_values(
        ["label_order", "scenario_order", "method_order", "window_order"]
    ).reset_index(drop=True)

    return df


def discover_profiles(profile_dir: Path) -> list[Path]:
    if not profile_dir.exists():
        raise FileNotFoundError(f"Feasibility profile directory not found: {profile_dir}")
    return sorted(profile_dir.glob("*.csv"))


def parse_profile_name(path: Path) -> dict:
    m = PROFILE_RE.match(path.name)
    if not m:
        raise ValueError(f"Unexpected feasibility profile filename: {path.name}")

    return {
        "method": m.group(1).upper(),
        "label": m.group(2).lower(),
        "date": m.group(3),
        "window_name": m.group(4),
        "scenario": m.group(5).lower(),
    }


def load_mechanism_profiles() -> dict[str, pd.DataFrame]:
    """
    Load two mechanism profiles:
    - SMA cloudy 2024-04-03 60min small
    - EMA cloudy 2024-04-03 60min small
    """
    profiles = discover_profiles(FEAS_PROFILE_DIR)
    selected = {}

    for p in profiles:
        meta = parse_profile_name(p)
        if (
            meta["label"] == MECH_LABEL
            and meta["date"] == MECH_DATE
            and meta["window_name"] == MECH_WINDOW
            and meta["scenario"] == MECH_SCENARIO
            and meta["method"] in {"SMA", "EMA"}
        ):
            df = pd.read_csv(p)
            df["timestamp"] = pd.to_datetime(df["timestamp"], errors="coerce")
            df = df.dropna(subset=["timestamp"]).sort_values("timestamp")
            df = df.set_index("timestamp")
            selected[meta["method"]] = df

    if "SMA" not in selected or "EMA" not in selected:
        raise FileNotFoundError(
            "Mechanism profiles not found for cloudy-2024-04-03-60min-small "
            "for both SMA and EMA."
        )

    return selected


# =========================
# 2) TABLE HELPERS
# =========================
def autosize_excel_columns(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 35)


def style_excel_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="000000")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    autosize_excel_columns(ws)


def save_html_table(df: pd.DataFrame, filepath: Path, title: str):
    styles = """
    <style>
        body { font-family: Calibri, Arial, sans-serif; margin: 24px; }
        h3 { margin-bottom: 12px; }
        table {
            border-collapse: collapse;
            width: 100%;
            font-size: 11pt;
        }
        th, td {
            border: 1px solid #000;
            padding: 6px 8px;
            text-align: center;
            vertical-align: middle;
        }
        th {
            background-color: #D9EAF7;
            font-weight: bold;
        }
        tr:nth-child(even) {
            background-color: #F8FBFD;
        }
    </style>
    """
    df_html = df.copy()
    for c in df_html.columns:
        df_html[c] = df_html[c].astype(str).str.replace("\n", "<br>", regex=False)

    html = (
        f"<html><head>{styles}</head><body>"
        f"<h3>{title}</h3>"
        f"{df_html.to_html(index=False, border=0, escape=False)}"
        f"</body></html>"
    )

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(html)


# =========================
# 3) BUILD TABLES
# =========================
def format_day_table(df: pd.DataFrame, label: str) -> pd.DataFrame:
    d = df[df["label"] == label].copy()
    d = d.sort_values(["scenario_order", "window_order", "method_order"])

    out = d[[
        "method", "window_name", "scenario",
        "ramp_p95_kw_target", "ramp_p95_kw_actual",
        "throughput_req_kwh", "throughput_act_kwh",
        "energy_mismatch_kwh", "power_limit_frac", "soc_limit_frac"
    ]].copy()

    out.columns = [
        "Method", "Window", "Scenario",
        "RampP95 target", "RampP95 actual",
        "Throughput_req", "Throughput_act",
        "Mismatch", "f_p", "f_soc"
    ]

    for c in ["RampP95 target", "RampP95 actual", "Throughput_req", "Throughput_act", "Mismatch", "f_p", "f_soc"]:
        out[c] = out[c].astype(float).round(3)

    return out


def build_synthesis_table(df: pd.DataFrame) -> pd.DataFrame:
    rows = []

    for label in ["clear", "medium", "cloudy"]:
        for scenario in ["small", "large"]:
            d = df[(df["label"] == label) & (df["scenario"] == scenario)].copy()
            if d.empty:
                continue

            best = d.loc[d["energy_mismatch_kwh"].idxmin()]
            worst = d.loc[d["energy_mismatch_kwh"].idxmax()]

            if worst["power_limit_frac"] > worst["soc_limit_frac"]:
                bottleneck = "Power limit"
            elif worst["soc_limit_frac"] > worst["power_limit_frac"]:
                bottleneck = "SoC limit"
            else:
                bottleneck = "Mixed"

            rows.append({
                "Day": f"{label} ({worst['date']})",
                "Scenario": scenario,
                "Best mismatch case": f"{best['method']} {best['window_name']}",
                "Best mismatch": round(best["energy_mismatch_kwh"], 3),
                "Worst mismatch case": f"{worst['method']} {worst['window_name']}",
                "Worst mismatch": round(worst["energy_mismatch_kwh"], 3),
                "Dominant bottleneck": bottleneck,
            })

    return pd.DataFrame(rows)


def build_mechanism_table(df: pd.DataFrame) -> pd.DataFrame:
    d = df[
        (df["label"] == MECH_LABEL)
        & (df["date"] == MECH_DATE)
        & (df["window_name"] == MECH_WINDOW)
        & (df["scenario"] == MECH_SCENARIO)
        & (df["method"].isin(["SMA", "EMA"]))
    ].copy()

    d = d.sort_values("method_order")

    out = d[[
        "method",
        "ramp_p95_kw_target",
        "ramp_p95_kw_actual",
        "throughput_req_kwh",
        "throughput_act_kwh",
        "energy_mismatch_kwh",
        "power_limit_frac",
        "soc_limit_frac",
    ]].copy()

    out.columns = [
        "Method",
        "RampP95 target",
        "RampP95 actual",
        "Throughput_req",
        "Throughput_act",
        "Mismatch",
        "f_p",
        "f_soc",
    ]

    for c in out.columns[1:]:
        out[c] = out[c].astype(float).round(3)

    return out


def save_all_tables(df: pd.DataFrame):
    clear_tbl = format_day_table(df, "clear")
    medium_tbl = format_day_table(df, "medium")
    cloudy_tbl = format_day_table(df, "cloudy")
    synth_tbl = build_synthesis_table(df)
    mech_tbl = build_mechanism_table(df)

    # CSV
    clear_tbl.to_csv(TAB_DIR / "table_4_5_clear.csv", index=False)
    medium_tbl.to_csv(TAB_DIR / "table_4_6_medium.csv", index=False)
    cloudy_tbl.to_csv(TAB_DIR / "table_4_7_cloudy.csv", index=False)
    synth_tbl.to_csv(TAB_DIR / "table_4_8_synthesis.csv", index=False)
    mech_tbl.to_csv(TAB_DIR / "table_4_9_mechanism.csv", index=False)

    # HTML
    save_html_table(clear_tbl, TAB_DIR / "table_4_5_clear.html", "Table 4.5. Level-2 results for clear day")
    save_html_table(medium_tbl, TAB_DIR / "table_4_6_medium.html", "Table 4.6. Level-2 results for medium day")
    save_html_table(cloudy_tbl, TAB_DIR / "table_4_7_cloudy.html", "Table 4.7. Level-2 results for cloudy day")
    save_html_table(synth_tbl, TAB_DIR / "table_4_8_synthesis.html", "Table 4.8. Cross-day synthesis of Level-2 feasibility")
    save_html_table(mech_tbl, TAB_DIR / "table_4_9_mechanism.html", "Table 4.9. Mechanism-case Level-2 metrics")

    # XLSX
    xlsx_path = TAB_DIR / "level2_tables_for_word.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        clear_tbl.to_excel(writer, sheet_name="Table_4_5_Clear", index=False)
        medium_tbl.to_excel(writer, sheet_name="Table_4_6_Medium", index=False)
        cloudy_tbl.to_excel(writer, sheet_name="Table_4_7_Cloudy", index=False)
        synth_tbl.to_excel(writer, sheet_name="Table_4_8_Synthesis", index=False)
        mech_tbl.to_excel(writer, sheet_name="Table_4_9_Mechanism", index=False)

    wb = load_workbook(xlsx_path)
    for ws in wb.worksheets:
        style_excel_sheet(ws)
    wb.save(xlsx_path)

    print("[OK] Saved all Level-2 tables to:", TAB_DIR)


# =========================
# 4) BUILD FIGURES
# =========================
def plot_figure_4_5_mismatch(df: pd.DataFrame):
    fig, axs = plt.subplots(1, 3, figsize=(14, 4.5), sharey=True)

    combos = [
        ("SMA", "small"),
        ("SMA", "large"),
        ("EMA", "small"),
        ("EMA", "large"),
    ]

    for ax, label in zip(axs, ["clear", "medium", "cloudy"]):
        dlabel = df[df["label"] == label].copy()

        for method, scenario in combos:
            g = dlabel[(dlabel["method"] == method) & (dlabel["scenario"] == scenario)].copy()
            g = g.sort_values("window_order")
            ax.plot(
                [int(w.replace("min", "")) for w in g["window_name"]],
                g["energy_mismatch_kwh"],
                marker=MARKERS_METHOD[method],
                linestyle=LINESTYLE_SCEN[scenario],
                linewidth=1.6,
                markersize=5,
                label=f"{method}-{scenario}",
            )

        ax.set_title(label.capitalize())
        ax.set_xlabel("Window (min)")
        ax.grid(True, alpha=0.3)

    axs[0].set_ylabel("Energy mismatch (kWh)")
    handles, labels = axs[0].get_legend_handles_labels()
    fig.legend(
        handles, labels,
        loc="upper right",
        bbox_to_anchor=(0.985, 0.985),
        ncol=4,
        fontsize=9,
        framealpha=0.9,
    )
    fig.tight_layout(rect=[0, 0, 1, 0.93])

    out = FIG_DIR / "fig_4_5_mismatch_by_day.png"
    fig.savefig(out, dpi=300)
    plt.close(fig)
    print("[OK] Saved:", out)


def plot_figure_4_6_tradeoff(df: pd.DataFrame):
    def make_tradeoff_figure(df_in: pd.DataFrame, scenario: str, out_name: str, title: str):
        fig, axs = plt.subplots(1, 3, figsize=(15, 4.8), sharex=False, sharey=False)

        method_colors = {
            "SMA": "tab:blue",
            "EMA": "tab:orange",
        }
        method_markers = {
            "SMA": "o",
            "EMA": "^",
        }

        for ax, label in zip(axs, ["clear", "medium", "cloudy"]):
            dlabel = df_in[(df_in["label"] == label) & (df_in["scenario"] == scenario)].copy()

            for method, g in dlabel.groupby("method"):
                g = g.sort_values("window_order")

                color = method_colors[method]
                marker = method_markers[method]

                # actual trajectory
                ax.plot(
                    g["ramp_p95_kw_actual"],
                    g["throughput_act_kwh"],
                    linestyle="-",
                    linewidth=1.6,
                    color=color,
                    alpha=0.95,
                )
                ax.scatter(
                    g["ramp_p95_kw_actual"],
                    g["throughput_act_kwh"],
                    marker=marker,
                    facecolors=color,
                    edgecolors=color,
                    s=60,
                    zorder=3,
                )

                # target points only
                ax.scatter(
                    g["ramp_p95_kw_target"],
                    g["throughput_req_kwh"],
                    marker=marker,
                    facecolors="none",
                    edgecolors=color,
                    s=60,
                    linewidths=1.3,
                    zorder=3,
                )

                # connect target -> actual
                for _, r in g.iterrows():
                    ax.plot(
                        [r["ramp_p95_kw_target"], r["ramp_p95_kw_actual"]],
                        [r["throughput_req_kwh"], r["throughput_act_kwh"]],
                        color="gray",
                        alpha=0.25,
                        linewidth=0.9,
                    )

                    # annotate actual points only
                    ax.annotate(
                        r["window_name"].replace("min", ""),
                        (r["ramp_p95_kw_actual"], r["throughput_act_kwh"]),
                        textcoords="offset points",
                        xytext=(4, 4),
                        fontsize=7,
                    )

            # autoscale per panel
            xvals = pd.concat(
                [dlabel["ramp_p95_kw_target"], dlabel["ramp_p95_kw_actual"]],
                ignore_index=True
            ).dropna()
            yvals = pd.concat(
                [dlabel["throughput_req_kwh"], dlabel["throughput_act_kwh"]],
                ignore_index=True
            ).dropna()

            if not xvals.empty:
                xmin, xmax = float(xvals.min()), float(xvals.max())
                xspan = xmax - xmin
                xpad = 0.08 * xspan if xspan > 0 else max(10.0, 0.1 * xmax)
                ax.set_xlim(max(0.0, xmin - xpad), xmax + xpad)

            if not yvals.empty:
                ymin, ymax = float(yvals.min()), float(yvals.max())
                yspan = ymax - ymin
                ypad = 0.08 * yspan if yspan > 0 else max(50.0, 0.1 * ymax)
                ax.set_ylim(max(0.0, ymin - ypad), ymax + ypad)

            ax.set_title(label.capitalize())
            ax.set_xlabel("RampP95 (kW / 5 min)")
            ax.grid(True, alpha=0.3)

        axs[0].set_ylabel("Battery throughput (kWh)")

        legend_handles = [
            Line2D([0], [0], color="tab:blue", lw=1.8, label="SMA"),
            Line2D([0], [0], color="tab:orange", lw=1.8, label="EMA"),
            Line2D([0], [0], marker="o", color="black", markerfacecolor="none",
                   linestyle="None", markersize=7, label="target"),
            Line2D([0], [0], marker="o", color="black", markerfacecolor="black",
                   linestyle="-", lw=1.4, markersize=7, label="actual"),
        ]

        fig.legend(
        handles=legend_handles,
        loc="upper right",
        bbox_to_anchor=(0.985, 0.985),
        ncol=4,
        fontsize=8.5,
        framealpha=0.9,
    )

        fig.suptitle(title, y=0.99)
        fig.tight_layout(rect=[0, 0, 1, 0.93])

        out = FIG_DIR / out_name
        fig.savefig(out, dpi=300)
        plt.close(fig)
        print("[OK] Saved:", out)

    make_tradeoff_figure(
        df_in=df,
        scenario="small",
        out_name="fig_4_6_small_target_vs_actual.png",
        title="Target vs actual trade-off — Small BESS",
    )

    make_tradeoff_figure(
        df_in=df,
        scenario="large",
        out_name="fig_4_7_large_target_vs_actual.png",
        title="Target vs actual trade-off — Large BESS",
    )


def plot_figure_4_8_mechanism(df_sum: pd.DataFrame):
    profs = load_mechanism_profiles()
    sma = profs["SMA"].copy()
    ema = profs["EMA"].copy()

    fig, axs = plt.subplots(3, 2, figsize=(13, 9), sharex="col")

    # get SoC bounds and battery power limit from summary if available
    dtmp = df_sum[
        (df_sum["label"] == MECH_LABEL)
        & (df_sum["date"] == MECH_DATE)
        & (df_sum["window_name"] == MECH_WINDOW)
        & (df_sum["scenario"] == MECH_SCENARIO)
    ]
    soc_min = float(dtmp["soc_min"].iloc[0]) if "soc_min" in dtmp.columns and not dtmp.empty else 0.20
    soc_max = float(dtmp["soc_max"].iloc[0]) if "soc_max" in dtmp.columns and not dtmp.empty else 0.80
    pbatt_max = float(dtmp["Pbatt_max_kw"].iloc[0]) if "Pbatt_max_kw" in dtmp.columns and not dtmp.empty else 500.0

    for col, (method, dfp) in enumerate([("SMA", sma), ("EMA", ema)]):
        mask = dfp["is_active"].astype(int) == 1
        d = dfp.loc[mask].copy()

        # row 1: Ppv / Psmooth / Pout_act
        axs[0, col].plot(d.index, d["P_pv_kw"], label="Ppv", linewidth=1.1)
        axs[0, col].plot(d.index, d["P_smooth_kw"], label="Psmooth", linewidth=1.4)
        axs[0, col].plot(d.index, d["P_out_act_kw"], label="Pout_act", linewidth=1.4)
        axs[0, col].set_title(method)
        axs[0, col].set_ylabel("Power (kW)")
        axs[0, col].grid(True, alpha=0.3)

       # row 2: Pbatt req / act
        axs[1, col].plot(d.index, d["P_batt_req_kw"], label="Pbatt_req", linewidth=1.2)
        axs[1, col].plot(d.index, d["P_batt_act_kw"], label="Pbatt_act", linewidth=1.2)
        axs[1, col].axhline(0, color="black", linewidth=0.8, alpha=0.6)

        # battery power limits
        axs[1, col].axhline(pbatt_max, color="purple", linestyle="--", linewidth=0.9, alpha=0.8, label="+Pmax")
        axs[1, col].axhline(-pbatt_max, color="purple", linestyle="--", linewidth=0.9, alpha=0.8, label="-Pmax")

        axs[1, col].set_ylim(
            min(axs[1, col].get_ylim()[0], -pbatt_max - 50),
            max(axs[1, col].get_ylim()[1], pbatt_max + 50),
        )

        yticks = list(axs[1, col].get_yticks())
        yticks.extend([pbatt_max, -pbatt_max])
        yticks = sorted(set(round(float(y), 6) for y in yticks))
        axs[1, col].set_yticks(yticks)

        axs[1, col].set_ylabel("Battery power (kW)")
        axs[1, col].grid(True, alpha=0.3)

        # row 3: SoC_actual
        axs[2, col].plot(d.index, d["SoC_actual"], label="SoC_actual", linewidth=1.4)
        axs[2, col].axhline(soc_min, color="red", linestyle="--", linewidth=0.9, alpha=0.8, label="SoC_min")
        axs[2, col].axhline(soc_max, color="green", linestyle="--", linewidth=0.9, alpha=0.8, label="SoC_max")
        axs[2, col].set_ylabel("SoC")
        axs[2, col].set_xlabel("Time (HH:MM)")
        axs[2, col].grid(True, alpha=0.3)

        for r in range(3):
            axs[r, col].xaxis.set_major_locator(mdates.HourLocator(interval=1))
            axs[r, col].xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))

    # compact legends
    for r in range(3):
        for c in range(2):
            axs[r, c].legend(fontsize=8, loc="upper left", framealpha=0.9)

    fig.suptitle("Mechanism case: cloudy – 60 min – small BESS", y=0.995)
    fig.tight_layout(rect=[0, 0, 1, 0.97])

    out = FIG_DIR / "fig_4_8_mechanism_cloudy_60min_small.png"
    fig.savefig(out, dpi=300)
    plt.close(fig)
    print("[OK] Saved:", out)


# =========================
# 5) MAIN
# =========================
def main():
    df = load_feasibility_summary(FEAS_SUMMARY_CSV)

    save_all_tables(df)
    plot_figure_4_5_mismatch(df)
    plot_figure_4_6_tradeoff(df)
    plot_figure_4_8_mechanism(df)

    print("\nAll Level-2 assets generated successfully.")
    print("Figures:", FIG_DIR)
    print("Tables :", TAB_DIR)


if __name__ == "__main__":
    main()