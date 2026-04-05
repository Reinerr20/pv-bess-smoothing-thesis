import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

INPUT_CSV = "outputs_baseline_compare/smoothing_results_baseline_compare.csv"
OUT_DIR = "outputs_baseline_compare/tables_word"
os.makedirs(OUT_DIR, exist_ok=True)

LABEL_ORDER = {"clear": 0, "medium": 1, "cloudy": 2}
METHOD_ORDER = {"SMA": 0, "EMA": 1}
WINDOW_ORDER = {"10min": 10, "20min": 20, "30min": 30, "60min": 60}


def load_data(path: str) -> pd.DataFrame:
    df = pd.read_csv(path)

    # pastikan kolom penting ada
    required = [
        "label", "date", "method", "window_name",
        "raw_ramp_p95_kw", "ramp_p95_kw", "ramp_reduction_pct", "throughput_kwh"
    ]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    df["label_order"] = df["label"].map(LABEL_ORDER)
    df["method_order"] = df["method"].map(METHOD_ORDER)
    df["window_order"] = df["window_name"].map(WINDOW_ORDER)

    df = df.sort_values(["label_order", "window_order", "method_order"]).reset_index(drop=True)
    return df


def format_day_table(df: pd.DataFrame, label: str) -> pd.DataFrame:
    d = df[df["label"] == label].copy()
    d = d.sort_values(["window_order", "method_order"])

    out = d[[
        "method", "window_name", "ramp_p95_kw", "ramp_reduction_pct", "throughput_kwh"
    ]].copy()

    out.columns = [
        "Method",
        "Window",
        "RampP95 target (kW/5 min)",
        "Ramp reduction (%)",
        "Throughput_req (kWh)"
    ]

    out["RampP95 target (kW/5 min)"] = out["RampP95 target (kW/5 min)"].round(2)
    out["Ramp reduction (%)"] = out["Ramp reduction (%)"].round(2)
    out["Throughput_req (kWh)"] = out["Throughput_req (kWh)"].round(2)

    return out


def build_synthesis_table(df: pd.DataFrame) -> pd.DataFrame:
    rows = []

    for label in ["clear", "medium", "cloudy"]:
        d = df[df["label"] == label].copy()
        if d.empty:
            continue

        raw_ramp = d["raw_ramp_p95_kw"].iloc[0]

        best_ramp = d.loc[d["ramp_p95_kw"].idxmin()]
        best_tp = d.loc[d["throughput_kwh"].idxmin()]

        rows.append({
            "Day": f"{label} ({best_ramp['date']})",
            "Raw RampP95 (kW/5 min)": round(raw_ramp, 2),
            "Lowest RampP95 candidate": f"{best_ramp['method']} {best_ramp['window_name']}",
            "Lowest RampP95 value": round(best_ramp["ramp_p95_kw"], 2),
            "Throughput at lowest RampP95 (kWh)": round(best_ramp["throughput_kwh"], 2),
            "Lowest throughput candidate": f"{best_tp['method']} {best_tp['window_name']}",
            "Lowest throughput value (kWh)": round(best_tp["throughput_kwh"], 2),
            "RampP95 at lowest throughput": round(best_tp["ramp_p95_kw"], 2),
        })

    return pd.DataFrame(rows)


def save_html_table(df: pd.DataFrame, filepath: str, title: str):
    styles = """
    <style>
        body { font-family: Calibri, Arial, sans-serif; margin: 24px; }
        h3 { margin-bottom: 12px; }
        table {
            border-collapse: collapse;
            width: 100%;
            font-size: 11pt;
        }
        th, td {
            border: 1px solid #000;
            padding: 6px 8px;
            text-align: center;
        }
        th {
            background-color: #D9EAF7;
            font-weight: bold;
        }
        tr:nth-child(even) {
            background-color: #F8FBFD;
        }
    </style>
    """
    html = f"<html><head>{styles}</head><body><h3>{title}</h3>{df.to_html(index=False, border=0)}</body></html>"

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(html)


def autosize_excel_columns(ws):
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                value = str(cell.value) if cell.value is not None else ""
                max_length = max(max_length, len(value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 35)


def style_excel_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="000000")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    autosize_excel_columns(ws)


def main():
    df = load_data(INPUT_CSV)

    # tabel per hari
    clear_tbl = format_day_table(df, "clear")
    medium_tbl = format_day_table(df, "medium")
    cloudy_tbl = format_day_table(df, "cloudy")

    # tabel sintesis
    synth_tbl = build_synthesis_table(df)

    # save CSV
    clear_tbl.to_csv(os.path.join(OUT_DIR, "table_4_1_clear.csv"), index=False)
    medium_tbl.to_csv(os.path.join(OUT_DIR, "table_4_2_medium.csv"), index=False)
    cloudy_tbl.to_csv(os.path.join(OUT_DIR, "table_4_3_cloudy.csv"), index=False)
    synth_tbl.to_csv(os.path.join(OUT_DIR, "table_4_4_synthesis.csv"), index=False)

    # save HTML
    save_html_table(clear_tbl, os.path.join(OUT_DIR, "table_4_1_clear.html"), "Table 4.1. Level-1 results for clear day")
    save_html_table(medium_tbl, os.path.join(OUT_DIR, "table_4_2_medium.html"), "Table 4.2. Level-1 results for medium day")
    save_html_table(cloudy_tbl, os.path.join(OUT_DIR, "table_4_3_cloudy.html"), "Table 4.3. Level-1 results for cloudy day")
    save_html_table(synth_tbl, os.path.join(OUT_DIR, "table_4_4_synthesis.html"), "Table 4.4. Cross-day quantitative synthesis of Level-1 results")

    # save XLSX
    xlsx_path = os.path.join(OUT_DIR, "level1_tables_for_word.xlsx")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        clear_tbl.to_excel(writer, sheet_name="Table_4_1_Clear", index=False)
        medium_tbl.to_excel(writer, sheet_name="Table_4_2_Medium", index=False)
        cloudy_tbl.to_excel(writer, sheet_name="Table_4_3_Cloudy", index=False)
        synth_tbl.to_excel(writer, sheet_name="Table_4_4_Synthesis", index=False)

    wb = load_workbook(xlsx_path)
    for ws in wb.worksheets:
        style_excel_sheet(ws)
    wb.save(xlsx_path)

    print("All Level-1 tables saved to:")
    print(OUT_DIR)


if __name__ == "__main__":
    main()