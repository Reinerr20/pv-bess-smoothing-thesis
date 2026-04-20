from __future__ import annotations

from pathlib import Path
from typing import Iterable
import math

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import matplotlib.dates as mdates

try:
    from openpyxl import Workbook
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.styles import Font, PatternFill
except Exception:  # pragma: no cover
    Workbook = None


# =========================================================
# CONFIG
# =========================================================
BASE_DIR = Path(__file__).resolve().parent


def _first_existing(*paths: Path) -> Path:
    for p in paths:
        if p.exists():
            return p
    return paths[0]


AGING_SUMMARY_CSV = _first_existing(
    BASE_DIR / "outputs_level2_aging" / "aging_summary.csv",
    BASE_DIR / "aging_summary.csv",
)
AGING_SUMMARY_WITH_FEAS_CSV = _first_existing(
    BASE_DIR / "outputs_level2_aging" / "aging_summary_with_feasibility.csv",
    BASE_DIR / "aging_summary_with_feasibility.csv",
)
AGING_CYCLES_DETAIL_CSV = _first_existing(
    BASE_DIR / "outputs_level2_aging" / "aging_cycles_detail.csv",
    BASE_DIR / "aging_cycles_detail.csv",
)
LEVEL2_FEASIBILITY_CSV = _first_existing(
    BASE_DIR / "outputs_level2_feasibility" / "level2_feasibility_results.csv",
    BASE_DIR / "level2_feasibility_results.csv",
)

# Explicit mechanism-case profiles provided by user
SMA_MECH_PROFILE = _first_existing(
    BASE_DIR / "outputs_level2_feasibility" / "feasibility_profiles" / "SMA_cloudy_2024-04-03_60min_small.csv",
    BASE_DIR / "SMA_cloudy_2024-04-03_60min_small.csv",
)
EMA_MECH_PROFILE = _first_existing(
    BASE_DIR / "outputs_level2_feasibility" / "feasibility_profiles" / "EMA_cloudy_2024-04-03_60min_small.csv",
    BASE_DIR / "EMA_cloudy_2024-04-03_60min_small.csv",
)

OUT_DIR = BASE_DIR / "outputs_level2_aging_assets"
OUT_DIR.mkdir(parents=True, exist_ok=True)

FIG_4_9 = OUT_DIR / "fig_4_9_aging_response_by_day.png"
FIG_4_10 = OUT_DIR / "fig_4_10_mechanism_aging_cloudy_60min_small.png"
FIG_4_11 = OUT_DIR / "fig_4_11_damage_vs_throughput.png"

TABLE_4_10_CSV = OUT_DIR / "table_4_10_mechanism.csv"
TABLE_4_11_CSV = OUT_DIR / "table_4_11_aging_synthesis.csv"
TABLES_XLSX = OUT_DIR / "aging_tables_for_word.xlsx"
CAPTIONS_TXT = OUT_DIR / "captions_4_4.txt"

DAY_ORDER = ["clear", "medium", "cloudy"]
METHOD_ORDER = ["SMA", "EMA"]
SCENARIO_ORDER = ["small", "large"]
WINDOW_ORDER = ["10min", "20min", "30min", "60min"]
WINDOW_TO_MIN = {"10min": 10, "20min": 20, "30min": 30, "60min": 60}

DAY_COLORS = {
    "clear": "#1f77b4",
    "medium": "#ff7f0e",
    "cloudy": "#2ca02c",
}
METHOD_MARKERS = {
    "SMA": "o",
    "EMA": "^",
}
SCENARIO_LABEL = {
    "small": "Small BESS",
    "large": "Large BESS",
}


# =========================================================
# HELPERS
# =========================================================
def _require_columns(df: pd.DataFrame, required: Iterable[str], name: str) -> None:
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"{name} is missing required columns: {missing}")


def _ordered_categorical(series: pd.Series, order: list[str]) -> pd.Categorical:
    return pd.Categorical(series, categories=order, ordered=True)


def load_merged_summary() -> pd.DataFrame:
    """
    Preferred source: aging_summary_with_feasibility.csv.
    Fallback: merge aging_summary.csv and level2_feasibility_results.csv.
    """
    if AGING_SUMMARY_WITH_FEAS_CSV.exists():
        df = pd.read_csv(AGING_SUMMARY_WITH_FEAS_CSV)
    else:
        aging = pd.read_csv(AGING_SUMMARY_CSV)
        feas = pd.read_csv(LEVEL2_FEASIBILITY_CSV)
        merge_keys = ["label", "date", "method", "window_name", "scenario"]
        df = aging.merge(feas, on=merge_keys, how="left", suffixes=("", "_feas"))

    required = [
        "label", "date", "method", "window_name", "scenario",
        "damage_index", "n_cycles_total", "soc_span_actual",
        "throughput_act_kwh", "energy_mismatch_kwh",
        "power_limit_frac", "soc_limit_frac",
        "ramp_p95_kw_target", "ramp_p95_kw_actual",
    ]
    _require_columns(df, required, "aging_summary_with_feasibility")

    # normalize
    df["label"] = df["label"].str.lower()
    df["method"] = df["method"].str.upper()
    df["scenario"] = df["scenario"].str.lower()
    df["window_name"] = df["window_name"].astype(str)
    df["window_min"] = df["window_name"].map(WINDOW_TO_MIN)
    df["label_ord"] = _ordered_categorical(df["label"], DAY_ORDER)
    df["method_ord"] = _ordered_categorical(df["method"], METHOD_ORDER)
    df["scenario_ord"] = _ordered_categorical(df["scenario"], SCENARIO_ORDER)
    df["window_ord"] = _ordered_categorical(df["window_name"], WINDOW_ORDER)

    return df.sort_values(["label_ord", "scenario_ord", "method_ord", "window_ord"]).reset_index(drop=True)


def load_cycles_detail() -> pd.DataFrame:
    df = pd.read_csv(AGING_CYCLES_DETAIL_CSV)
    required = ["label", "date", "method", "window_name", "scenario", "dod_mid", "cycle_count"]
    _require_columns(df, required, "aging_cycles_detail")
    df["label"] = df["label"].str.lower()
    df["method"] = df["method"].str.upper()
    df["scenario"] = df["scenario"].str.lower()
    return df


def load_mechanism_profile(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Mechanism profile not found: {path}")

    df = pd.read_csv(path)
    required = ["timestamp", "SoC_actual", "is_active"]
    _require_columns(df, required, path.name)

    df["timestamp"] = pd.to_datetime(df["timestamp"], errors="coerce")
    df = df.dropna(subset=["timestamp"]).copy()
    df["SoC_actual"] = pd.to_numeric(df["SoC_actual"], errors="coerce")
    df["is_active"] = pd.to_numeric(df["is_active"], errors="coerce").fillna(0).astype(int)
    return df


def human_case(row: pd.Series) -> str:
    return f"{row['method']} {row['window_name']} {row['scenario']}"


def human_day(label: str, date_str: str) -> str:
    return f"{label} ({date_str})"


def save_excel_tables(table_410: pd.DataFrame, table_411: pd.DataFrame) -> None:
    if Workbook is None:
        return

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Table_4_10_Mechanism"
    ws2 = wb.create_sheet("Table_4_11_Synthesis")

    for ws, df in [(ws1, table_410), (ws2, table_411)]:
        for row in dataframe_to_rows(df, index=False, header=True):
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for col_cells in ws.columns:
            length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 24)

    wb.save(TABLES_XLSX)


def save_captions() -> None:
    text = """Gambar 4.9 Pengaruh horizon smoothing terhadap damage index relatif (baris atas) dan jumlah siklus total hasil rainflow counting (baris bawah) pada hari clear, medium, dan cloudy untuk skenario small dan large BESS.

Gambar 4.10 Perbandingan profil SoC_actual dan distribusi cycle count terhadap DoD bin untuk SMA dan EMA pada hari cloudy, horizon 60 menit, dan skenario small BESS.

Gambar 4.11 Hubungan antara throughput aktual baterai dan damage index relatif pada seluruh kombinasi metode, horizon smoothing, hari representatif, dan skenario BESS.

Tabel 4.10 Ringkasan metrik feasibility dan aging untuk kasus representatif mekanisme pada hari cloudy, horizon 60 menit, dan skenario small BESS.

Tabel 4.11 Ringkasan pola utama aging proxy lintas hari representatif.
"""
    CAPTIONS_TXT.write_text(text, encoding="utf-8")


# =========================================================
# FIGURE 4.9
# =========================================================
def make_fig_4_9(df: pd.DataFrame) -> None:
    fig, axes = plt.subplots(2, 3, figsize=(14, 7), sharex="col")
    linestyle_map = {"small": "-", "large": "--"}
    color_map = {"SMA": "#1f77b4", "EMA": "#ff7f0e"}

    for col, day in enumerate(DAY_ORDER):
        sub = df[df["label"] == day].copy()
        ax_top = axes[0, col]
        ax_bot = axes[1, col]

        for method in METHOD_ORDER:
            for scenario in SCENARIO_ORDER:
                g = sub[(sub["method"] == method) & (sub["scenario"] == scenario)].copy()
                if g.empty:
                    continue
                g = g.sort_values("window_min")
                label = f"{method}-{scenario}"
                ax_top.plot(
                    g["window_min"], g["damage_index"],
                    marker=METHOD_MARKERS[method],
                    linestyle=linestyle_map[scenario],
                    color=color_map[method],
                    linewidth=1.8,
                    markersize=6,
                    label=label,
                )
                ax_bot.plot(
                    g["window_min"], g["n_cycles_total"],
                    marker=METHOD_MARKERS[method],
                    linestyle=linestyle_map[scenario],
                    color=color_map[method],
                    linewidth=1.8,
                    markersize=6,
                    label=label,
                )

        ax_top.set_title(day.capitalize())
        ax_top.grid(True, alpha=0.25)
        ax_bot.grid(True, alpha=0.25)
        ax_bot.set_xticks([10, 20, 30, 60])
        ax_bot.set_xticklabels(["10", "20", "30", "60"])
        ax_bot.set_xlabel("Horizon smoothing (min)")

        if col == 0:
            ax_top.set_ylabel("Damage index")
            ax_bot.set_ylabel("Total cycles")
        else:
            ax_top.set_ylabel("")
            ax_bot.set_ylabel("")

    handles = [
        Line2D([0], [0], color=color_map["SMA"], marker=METHOD_MARKERS["SMA"], linestyle="-", linewidth=1.8, label="SMA-small"),
        Line2D([0], [0], color=color_map["SMA"], marker=METHOD_MARKERS["SMA"], linestyle="--", linewidth=1.8, label="SMA-large"),
        Line2D([0], [0], color=color_map["EMA"], marker=METHOD_MARKERS["EMA"], linestyle="-", linewidth=1.8, label="EMA-small"),
        Line2D([0], [0], color=color_map["EMA"], marker=METHOD_MARKERS["EMA"], linestyle="--", linewidth=1.8, label="EMA-large"),
    ]
    fig.legend(handles=handles, loc="lower center", ncol=4, frameon=False, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(FIG_4_9, dpi=300, bbox_inches="tight")
    plt.close(fig)


# =========================================================
# TABLE 4.10 + FIGURE 4.10
# =========================================================
def build_table_4_10(df: pd.DataFrame) -> pd.DataFrame:
    mech = df[
        (df["label"] == "cloudy") &
        (df["date"] == "2024-04-03") &
        (df["window_name"] == "60min") &
        (df["scenario"] == "small") &
        (df["method"].isin(["SMA", "EMA"]))
    ].copy()

    mech = mech.sort_values("method_ord")

    out = mech[[
        "method",
        "throughput_act_kwh",
        "energy_mismatch_kwh",
        "power_limit_frac",
        "soc_limit_frac",
        "soc_span_actual",
        "n_cycles_total",
        "damage_index",
    ]].copy()

    out.columns = [
        "Method",
        "Throughput_act",
        "Mismatch",
        "fp",
        "fsoc",
        "soc_span_actual",
        "n_cycles_total",
        "damage_index",
    ]

    round_map = {
        "Throughput_act": 3,
        "Mismatch": 3,
        "fp": 3,
        "fsoc": 3,
        "soc_span_actual": 3,
        "n_cycles_total": 1,
        "damage_index": 6,
    }
    for c, nd in round_map.items():
        out[c] = out[c].astype(float).round(nd)

    out.to_csv(TABLE_4_10_CSV, index=False)
    return out


def make_fig_4_10(df: pd.DataFrame, cycles_detail: pd.DataFrame) -> None:
    sma_prof = load_mechanism_profile(SMA_MECH_PROFILE)
    ema_prof = load_mechanism_profile(EMA_MECH_PROFILE)

    sma_prof = sma_prof[sma_prof["is_active"] == 1].copy()
    ema_prof = ema_prof[ema_prof["is_active"] == 1].copy()

    # read SoC bounds from merged summary (same for both mechanism rows)
    mech = df[
        (df["label"] == "cloudy") &
        (df["date"] == "2024-04-03") &
        (df["window_name"] == "60min") &
        (df["scenario"] == "small")
    ].copy()
    soc_min = float(mech["soc_min"].iloc[0]) if "soc_min" in mech.columns else 0.20
    soc_max = float(mech["soc_max"].iloc[0]) if "soc_max" in mech.columns else 0.80

    cyc = cycles_detail[
        (cycles_detail["label"] == "cloudy") &
        (cycles_detail["date"] == "2024-04-03") &
        (cycles_detail["window_name"] == "60min") &
        (cycles_detail["scenario"] == "small")
    ].copy()

    sma_cyc = cyc[cyc["method"] == "SMA"].copy()
    ema_cyc = cyc[cyc["method"] == "EMA"].copy()

    dod_all = sorted(set(sma_cyc["dod_mid"].tolist()) | set(ema_cyc["dod_mid"].tolist()))
    dod = np.array(dod_all, dtype=float)

    def aligned_counts(g: pd.DataFrame) -> np.ndarray:
        mapper = dict(zip(g["dod_mid"], g["cycle_count"]))
        return np.array([mapper.get(x, 0.0) for x in dod], dtype=float)

    sma_counts = aligned_counts(sma_cyc)
    ema_counts = aligned_counts(ema_cyc)

    fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))

    # (a) SoC profile
    ax = axes[0]
    ax.plot(sma_prof["timestamp"], sma_prof["SoC_actual"], label="SMA", linewidth=1.9)
    ax.plot(ema_prof["timestamp"], ema_prof["SoC_actual"], label="EMA", linewidth=1.9)
    ax.axhline(soc_min, linestyle="--", linewidth=1.2, color="gray", label="SoC min/max")
    ax.axhline(soc_max, linestyle="--", linewidth=1.2, color="gray")

    ax.set_title("(a) SoC_actual during active PV window")
    ax.set_xlabel("Time (HH:MM)")
    ax.set_ylabel("SoC")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False)

    # x-axis: jam saja, tanpa tanggal
    ax.xaxis.set_major_locator(mdates.HourLocator(interval=2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
    ax.tick_params(axis="x", rotation=0)

    # (b) DoD distribution
    ax = axes[1]
    width = 0.018
    ax.bar(dod - width/2, sma_counts, width=width, label="SMA", alpha=0.85)
    ax.bar(dod + width/2, ema_counts, width=width, label="EMA", alpha=0.85)
    ax.set_title("(b) Cycle-count distribution by DoD bin")
    ax.set_xlabel("DoD bin midpoint")
    ax.set_ylabel("Cycle count")
    ax.grid(True, axis="y", alpha=0.25)
    ax.legend(frameon=False)
    ax.set_xlim(left=max(0, dod.min() - 0.03), right=min(1.0, dod.max() + 0.05))

    fig.tight_layout()
    fig.savefig(FIG_4_10, dpi=300, bbox_inches="tight")
    plt.close(fig)


# =========================================================
# FIGURE 4.11 + TABLE 4.11
# =========================================================
def make_fig_4_11(df: pd.DataFrame) -> None:
    fig, axes = plt.subplots(1, 3, figsize=(14.5, 4.8), sharey=True)

    color_map = {
        "SMA": "#1f77b4",
        "EMA": "#ff7f0e",
    }

    for i, day in enumerate(DAY_ORDER):
        ax = axes[i]
        sub = df[df["label"] == day].copy()

        for method in METHOD_ORDER:
            for scenario in SCENARIO_ORDER:
                g = sub[(sub["method"] == method) & (sub["scenario"] == scenario)].copy()
                if g.empty:
                    continue

                g = g.sort_values("window_min")

                color = color_map[method]
                marker = METHOD_MARKERS[method]
                line_style = "--" if scenario == "small" else "-"

                # connector line
                ax.plot(
                    g["throughput_act_kwh"],
                    g["damage_index"],
                    color=color,
                    linestyle=line_style,
                    linewidth=1.2,
                    alpha=0.45,
                    zorder=1,
                )

                # marker: hollow untuk small, filled untuk large
                if scenario == "small":
                    face = "none"
                    lw = 1.8
                else:
                    face = color
                    lw = 0.8

                ax.scatter(
                    g["throughput_act_kwh"],
                    g["damage_index"],
                    marker=marker,
                    s=75,
                    facecolors=face,
                    edgecolors=color,
                    linewidths=lw,
                    alpha=0.95,
                    zorder=2,
                )

        ax.set_title(day.capitalize())
        ax.set_xlabel("Throughput_act (kWh)")
        ax.grid(True, alpha=0.25)

        if i == 0:
            ax.set_ylabel("Damage index")
        else:
            ax.set_ylabel("")

    # legend method
    method_handles = [
        Line2D(
            [0], [0],
            marker="o",
            linestyle="",
            color=color_map["SMA"],
            markerfacecolor=color_map["SMA"],
            markersize=8,
            label="SMA",
        ),
        Line2D(
            [0], [0],
            marker="^",
            linestyle="",
            color=color_map["EMA"],
            markerfacecolor=color_map["EMA"],
            markersize=8,
            label="EMA",
        ),
    ]

    # legend scenario
    scenario_handles = [
        Line2D(
            [0], [0],
            marker="o",
            linestyle="--",
            color="black",
            markerfacecolor="none",
            markersize=8,
            label="Small BESS",
        ),
        Line2D(
            [0], [0],
            marker="o",
            linestyle="-",
            color="black",
            markerfacecolor="black",
            markersize=8,
            label="Large BESS",
        ),
    ]

    leg1 = fig.legend(
        handles=method_handles,
        loc="lower center",
        bbox_to_anchor=(0.34, -0.065),
        ncol=2,
        frameon=False,
        title="Method",
        title_fontsize=13,
        fontsize=11,
        handlelength=2.0,
        columnspacing=1.6,
        labelspacing=0.8,
        borderpad=0.6,
    )
    leg1.get_title().set_fontweight("bold")
    fig.add_artist(leg1)

    leg2 = fig.legend(
        handles=scenario_handles,
        loc="lower center",
        bbox_to_anchor=(0.77, -0.065),
        ncol=2,
        frameon=False,
        title="Scenario",
        title_fontsize=13,
        fontsize=11,
        handlelength=2.0,
        columnspacing=1.6,
        labelspacing=0.8,
        borderpad=0.6,
    )
    leg2.get_title().set_fontweight("bold")

    fig.suptitle("Damage index vs actual battery throughput by representative day", y=1.03)
    fig.tight_layout(rect=(0, 0.12, 1, 1))
    fig.savefig(FIG_4_11, dpi=300, bbox_inches="tight")
    plt.close(fig)


def build_table_4_11(df: pd.DataFrame) -> pd.DataFrame:
    rows = []

    mean_by_day = df.groupby("label", as_index=False)["damage_index"].mean()
    lowest_day = mean_by_day.sort_values("damage_index").iloc[0]["label"]
    highest_day = mean_by_day.sort_values("damage_index").iloc[-1]["label"]

    for day in DAY_ORDER:
        sub = df[df["label"] == day].copy()
        sub = sub.sort_values(["damage_index", "window_min", "scenario_ord", "method_ord"])
        min_row = sub.iloc[0]
        max_row = sub.iloc[-1]

        # safe, concise pattern
        if day == lowest_day:
            pattern = "Damage relatif paling rendah; horizon besar cenderung menurunkan damage."
        elif day == highest_day:
            pattern = "Damage relatif paling tinggi; constraint paling terasa dan damage tidak identik dengan throughput."
        else:
            pattern = "Damage cenderung turun saat horizon membesar; pengaruh metode tetap bergantung konteks."

        rows.append({
            "Hari": human_day(min_row["label"], min_row["date"]),
            "Kandidat damage minimum": human_case(min_row),
            "Nilai damage minimum": round(float(min_row["damage_index"]), 6),
            "Kandidat damage maksimum": human_case(max_row),
            "Nilai damage maksimum": round(float(max_row["damage_index"]), 6),
            "Pola utama": pattern,
        })

    out = pd.DataFrame(rows)
    out.to_csv(TABLE_4_11_CSV, index=False)
    return out


# =========================================================
# MAIN
# =========================================================
def main() -> None:
    df = load_merged_summary()
    cycles_detail = load_cycles_detail()

    # Build tables first (also useful for manuscript drafting)
    table_410 = build_table_4_10(df)
    table_411 = build_table_4_11(df)

    # Build figures
    make_fig_4_9(df)
    make_fig_4_10(df, cycles_detail)
    make_fig_4_11(df)

    # Companion outputs
    save_excel_tables(table_410, table_411)
    save_captions()

    print("Aging assets generated successfully:")
    print(f"- {FIG_4_9}")
    print(f"- {FIG_4_10}")
    print(f"- {FIG_4_11}")
    print(f"- {TABLE_4_10_CSV}")
    print(f"- {TABLE_4_11_CSV}")
    if TABLES_XLSX.exists():
        print(f"- {TABLES_XLSX}")
    print(f"- {CAPTIONS_TXT}")


if __name__ == "__main__":
    main()